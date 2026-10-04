import tempfile
import unittest
from pathlib import Path

import openpyxl

from tdoc_converter.tdoc_list import available_meetings, find_tdoc_list, load_tdoc_list, meeting_folder_candidates, select_agenda_item

HEADER = ["TDoc", "Title", "Source", "Agenda item", "TDoc Status", "Uploaded"]


class TdocListTest(unittest.TestCase):
    def make_xlsx(self, rows):
        d = Path(tempfile.mkdtemp())
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "TDoc_List"
        ws.append(HEADER)
        for r in rows:
            ws.append(r)
        p = d / "TDoc_List_Meeting_RAN1#126-bis.xlsx"
        wb.save(p)
        return p

    def test_select_exact_agenda_item(self):
        p = self.make_xlsx([
            ["R1-1", "a", "X", "10.5.2.2", "available", "2026-10-01"],
            ["R1-2", "b", "Y", "10.5.2", "available", "2026-10-01"],
            ["R1-3", "c", "Z", "10.5.2.2", "reserved", None],
            ["R1-4", "d", "W", "10.5.2.20", "available", "2026-10-01"],
        ])
        got = select_agenda_item(load_tdoc_list(p), "10.5.2.2")
        self.assertEqual([t.number for t in got], ["R1-1", "R1-3"])
        self.assertEqual([t.downloadable for t in got], [True, False])

    def test_download_url_comes_from_cell_hyperlink(self):
        d = Path(tempfile.mkdtemp())
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "TDoc_List"
        ws.append(HEADER)
        ws.append(["R1-1", "a", "X", "10.5.2.2", "available", "2026-10-01"])
        ws.append(["R1-2", "b", "Y", "10.5.2.2", "reserved", None])
        ws["A2"].hyperlink = "https://www.3gpp.org/ftp/tsg_ran/WG1_RL1/TSGR1_126b/Docs/R1-1.zip"
        p = d / "list.xlsx"
        wb.save(p)
        a, b = select_agenda_item(load_tdoc_list(p), "10.5.2.2")
        self.assertEqual(a.url, "https://www.3gpp.org/ftp/tsg_ran/WG1_RL1/TSGR1_126b/Docs/R1-1.zip")
        self.assertEqual(b.url, "")

    def test_find_tdoc_list_by_meeting(self):
        d = Path(tempfile.mkdtemp())
        for name in ("TDoc_List_Meeting_RAN1#126-bis.xlsx", "TDoc_List_Meeting_RAN1#126.xlsx", "notes.xlsx"):
            (d / name).touch()
        self.assertEqual(len(available_meetings(d)), 2)
        for spec in ("126bis", "126-bis", "126b", "RAN1#126-bis", "#126bis"):
            self.assertEqual(find_tdoc_list(d, spec).name, "TDoc_List_Meeting_RAN1#126-bis.xlsx", spec)
        self.assertEqual(find_tdoc_list(d, "126").name, "TDoc_List_Meeting_RAN1#126.xlsx")
        with self.assertRaisesRegex(ValueError, "choose one with --meeting"):
            find_tdoc_list(d)
        with self.assertRaisesRegex(ValueError, "available"):
            find_tdoc_list(d, "127")
        (d / "TDoc_List_Meeting_RAN1#126.xlsx").unlink()
        self.assertEqual(find_tdoc_list(d).name, "TDoc_List_Meeting_RAN1#126-bis.xlsx")

    def test_meeting_folder_candidates(self):
        self.assertEqual(meeting_folder_candidates("TDoc_List_Meeting_RAN1#126-bis.xlsx")[0], "TSGR1_126b")
        self.assertEqual(meeting_folder_candidates("TDoc_List_Meeting_RAN1#126.xlsx"), ["TSGR1_126"])
        self.assertIn("TSGR1_126bis", meeting_folder_candidates("TDoc_List_Meeting_RAN1#126-bis.xlsx"))
        with self.assertRaises(ValueError):
            meeting_folder_candidates("something.xlsx")


if __name__ == "__main__":
    unittest.main()
