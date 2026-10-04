import base64
import io
import struct
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from PIL import Image

from docx import Document
from docx.oxml.ns import qn
from pptx import Presentation

from tdoc_converter import convert, fix_images
from tdoc_converter.convert import to_markdown
from tdoc_converter.extract import extract_statements


def make_docx_with_label_numbering(path: Path):
    """Docx whose 'List Number' style renders as 'Proposal %1:' (label absent from the text)."""
    doc = Document()
    numbering = doc.part.numbering_part.element
    style_num = doc.styles["List Number"].element.pPr.numPr.numId.val
    abs_id = [n for n in numbering.findall(qn("w:num")) if n.get(qn("w:numId")) == str(style_num)][0] \
        .find(qn("w:abstractNumId")).get(qn("w:val"))
    for a in numbering.findall(qn("w:abstractNum")):
        if a.get(qn("w:abstractNumId")) == abs_id:
            a.find(qn("w:lvl")).find(qn("w:lvlText")).set(qn("w:val"), "Proposal %1:")
    doc.add_heading("Intro", 1)
    doc.add_paragraph("Support single TB over 400 MHz.", style="List Number")
    doc.add_paragraph("Some analysis.")
    doc.add_paragraph("Prefer layer-common bundling.", style="List Number")
    doc.add_paragraph("Observation 1: Plain-text label survives.")
    doc.save(path)


class ConvertTest(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())

    def test_auto_numbered_labels_are_recovered(self):
        p = self.tmp / "a.docx"
        make_docx_with_label_numbering(p)
        got = {s.label: s.text for s in extract_statements(to_markdown(p))}
        self.assertEqual(got, {
            "Proposal 1": "Support single TB over 400 MHz.",
            "Proposal 2": "Prefer layer-common bundling.",
            "Observation 1": "Plain-text label survives.",
        })

    def test_pptx(self):
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Summary"
        slide.placeholders[1].text = "Proposal 1: Support X"
        p = self.tmp / "a.pptx"
        prs.save(p)
        (s,) = extract_statements(to_markdown(p))
        self.assertEqual((s.label, s.text), ("Proposal 1", "Support X"))

    def png_bytes(self, color):
        buf = io.BytesIO()
        Image.new("RGB", (8, 8), color).save(buf, "PNG")
        return buf.getvalue()

    def test_docx_images_are_saved_and_linked(self):
        red, blue = self.png_bytes("red"), self.png_bytes("blue")
        doc = Document()
        doc.add_paragraph("Proposal 1: See figure.")
        for data in (red, blue, red):                      # red appears twice
            doc.add_picture(io.BytesIO(data))
        p = self.tmp / "R1-1.docx"
        doc.save(p)
        md_path = convert.convert_to_file(p, self.tmp / "markdown" / "R1-1.md")
        text = md_path.read_text()
        links = [l for l in text.split() if l.startswith("![")]
        self.assertEqual(links, ["![image](images/R1-1/img001.png)", "![image](images/R1-1/img002.png)",
                                 "![image](images/R1-1/img001.png)"])   # duplicate reuses the file
        img_dir = md_path.parent / "images" / "R1-1"
        self.assertEqual(sorted(f.name for f in img_dir.iterdir()), ["img001.png", "img002.png"])
        self.assertEqual((img_dir / "img001.png").read_bytes(), red)
        self.assertNotIn("base64", text)
        # the extractor still sees the statement
        self.assertEqual([s.text for s in extract_statements(text)], ["See figure."])

    def test_images_dropped_without_image_dir(self):
        doc = Document()
        doc.add_picture(io.BytesIO(self.png_bytes("red")))
        p = self.tmp / "a.docx"
        doc.save(p)
        self.assertNotIn("![", to_markdown(p))

    def test_emf_kept_when_no_converter(self):
        b64 = base64.b64encode(b"EMF-DATA").decode()
        text = f"![](data:image/x-emf;base64,{b64})"
        with mock.patch("shutil.which", return_value=None), mock.patch.object(convert, "_soffice", return_value=None):
            out = convert.externalize_images(text, self.tmp / "im", "images/x/")
        self.assertEqual(out, "![image](images/x/img001.emf)")
        self.assertEqual((self.tmp / "im" / "img001.emf").read_bytes(), b"EMF-DATA")

    def fake_soffice(self, drawing):
        """Stand-in for `soffice --convert-to png --outdir DIR file`: writes a page with `drawing` on it."""
        def run(cmd, **kw):
            page = Image.new("RGB", (200, 300), "white")
            if drawing:
                page.paste((255, 0, 0), (50, 100, 90, 130))
            page.save(Path(cmd[cmd.index("--outdir") + 1]) / (Path(cmd[-1]).stem + ".png"))
        return run

    def test_emf_rendered_with_libreoffice_and_cropped(self):
        emf = self.tmp / "a.emf"
        emf.write_bytes(b"EMF-DATA")
        with mock.patch.object(convert, "_soffice", return_value="/bin/soffice"), \
                mock.patch("subprocess.run", self.fake_soffice(True)):
            png = convert._vector_to_png(emf)
        self.assertEqual(png, self.tmp / "a.png")
        with Image.open(png) as im:
            self.assertEqual(im.size, (40 + 24, 30 + 24))      # figure + 12px margin on each side

    def test_blank_render_without_bitmap_keeps_emf(self):
        emf = self.tmp / "a.emf"
        emf.write_bytes(b"EMF-DATA")
        with mock.patch.object(convert, "_soffice", return_value="/bin/soffice"), \
                mock.patch("subprocess.run", self.fake_soffice(False)):
            self.assertIsNone(convert._vector_to_png(emf))
        self.assertFalse((self.tmp / "a.png").exists())

    def test_blank_render_falls_back_to_emfplus_bitmap(self):
        pixels = bytes([0, 0, 255, 255]) * 6                   # BGRA red, 3x2
        image_obj = struct.pack("<IIiiiII", 0xDBC01002, 1, 3, 2, 12, 0x0026200A, 0) + pixels
        record = struct.pack("<HHII", 0x4008, 0x0500, 12 + len(image_obj), len(image_obj)) + image_obj
        comment = struct.pack("<II", 70, 0) + struct.pack("<I", 4 + len(record)) + b"EMF+" + record
        comment = struct.pack("<II", 70, 12 + 4 + len(record)) + comment[8:]
        emf = self.tmp / "a.emf"
        emf.write_bytes(struct.pack("<II", 1, 8) + comment)
        with mock.patch.object(convert, "_soffice", return_value="/bin/soffice"), \
                mock.patch("subprocess.run", self.fake_soffice(False)):
            png = convert._vector_to_png(emf)
        with Image.open(png) as im:
            self.assertEqual((im.size, im.convert("RGB").getpixel((0, 0))), ((3, 2), (255, 0, 0)))

    def test_soffice_found_in_mac_app_bundle(self):
        with mock.patch("shutil.which", return_value=None), \
                mock.patch.object(Path, "exists", return_value=True):
            self.assertEqual(convert._soffice(), convert._MAC_SOFFICE)

    def test_fix_markdown_dir_converts_and_relinks(self):
        img = self.tmp / "images" / "R1-1"
        img.mkdir(parents=True)
        (img / "img001.emf").write_bytes(b"EMF-DATA")
        (self.tmp / "R1-1.md").write_text("![image](images/R1-1/img001.emf)\n")
        with mock.patch.object(convert, "_soffice", return_value="/bin/soffice"), \
                mock.patch("subprocess.run", self.fake_soffice(True)):
            self.assertEqual(fix_images.fix_markdown_dir(self.tmp), (1, []))
        self.assertEqual((self.tmp / "R1-1.md").read_text(), "![image](images/R1-1/img001.png)\n")
        self.assertEqual(sorted(p.name for p in img.iterdir()), ["img001.png"])

    def test_unsupported_extension(self):
        with self.assertRaises(ValueError):
            to_markdown(self.tmp / "a.pdf")


if __name__ == "__main__":
    unittest.main()
